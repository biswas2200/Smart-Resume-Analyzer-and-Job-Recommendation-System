"""DOCX implementation of DocumentTextExtractor, via python-docx. Same
impure/pure split as pdf_extractor.py: DocxTextExtractor.extract opens the
file, _document_lines walks it, and _lines_to_text (pure) turns the walked
lines into text so the assembly rules can be tested without a .docx file.

Reading order: body paragraphs and tables in document order (a resume's
two-column layouts are usually tables -- each row's cells are emitted left to
right, one line per cell, so a skills sidebar ends up next to, not inside,
the content it sits beside), then every header and footer once. python-docx
exposes paragraphs and tables as separate collections, so document order is
recovered by walking the body element's children directly.

Hidden text, as in pdf_extractor.py: a run a reader can't see is left out
and counted, so white-on-white keywords or instructions planted for a parser
never reach it. A run is hidden when it is marked vanished (Word's "Hidden"
font option), set under _MIN_VISIBLE_FONT_SIZE points, or coloured white or
nearly white with no dark shading or highlight behind it (white text in a
dark table cell or header bar is visible and kept). Formatting is read from
the run itself, then its character style, then its paragraph's style chain;
document defaults and themes beyond white theme colours are not resolved.
"""

import io
from collections.abc import Iterable

from app.services.document_extraction.base import ExtractedText

# Same limits as pdf_extractor.py.
_MIN_VISIBLE_FONT_SIZE = 3.0
_NEAR_WHITE_TOLERANCE = 0.05

# Theme colours that are white in Word's default themes ("Background 1",
# "Light 1"); a run coloured with one of these reads as white on a white page.
_WHITE_THEME_COLORS = frozenset({"BACKGROUND_1", "LIGHT_1"})


class DocxTextExtractor:
    """DocumentTextExtractor for .docx (Office Open XML) files."""

    def extract(self, file_bytes: bytes) -> ExtractedText:
        # Imported lazily so python-docx only has to be importable when a DOCX
        # is actually uploaded, matching pdf_extractor's pattern.
        from docx import Document

        document = Document(io.BytesIO(file_bytes))
        lines, hidden_word_count = [], 0
        for text, hidden in _document_lines(document):
            lines.append(text)
            hidden_word_count += hidden
        return ExtractedText(_lines_to_text(lines), hidden_word_count)


def _document_lines(document) -> Iterable[tuple[str, int]]:
    """(visible text, hidden word count) for each paragraph, in reading order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield _paragraph_text(Paragraph(child, document))
        elif tag == "tbl":
            for row in Table(child, document).rows:
                seen_cells = set()
                for cell in row.cells:
                    # Merged cells repeat the same underlying element.
                    if id(cell._tc) in seen_cells:
                        continue
                    seen_cells.add(id(cell._tc))
                    cell_fill = _shading_fill(cell._tc.tcPr)
                    for paragraph in cell.paragraphs:
                        yield _paragraph_text(paragraph, cell_fill)

    for section in document.sections:
        for part in (section.header, section.footer):
            if part is None or part.is_linked_to_previous:
                continue
            for paragraph in part.paragraphs:
                yield _paragraph_text(paragraph)


def _paragraph_text(paragraph, cell_fill: str | None = None) -> tuple[str, int]:
    """The paragraph's visible text and how many words its hidden runs held.
    Walks every run, including those inside hyperlinks and tracked insertions,
    in document order. Runs of a text box anchored in the paragraph belong to
    the text box's own paragraphs, so they are skipped, as paragraph.text does."""
    from docx.oxml.ns import qn
    from docx.text.run import Run

    p_tag = qn("w:p")
    paragraph_fill = _shading_fill(paragraph._p.pPr)
    visible, hidden = [], 0
    for element in paragraph._p.iter(qn("w:r")):
        if next(element.iterancestors(p_tag), None) is not paragraph._p:
            continue
        run = Run(element, paragraph)
        text = run.text
        if not text:
            continue
        background = _shading_fill(element.rPr) or _highlight(element.rPr) or paragraph_fill or cell_fill
        if _is_hidden_run(
            vanished=_resolve(run, paragraph, lambda font: font.hidden),
            size_pt=_resolve(run, paragraph, lambda font: font.size.pt if font.size is not None else None),
            color=_resolve(run, paragraph, _font_color),
            background=background,
        ):
            hidden += len(text.split())
        else:
            visible.append(text)
    return "".join(visible), hidden


def _is_hidden_run(vanished: bool | None, size_pt: float | None, color: str | None, background: str | None) -> bool:
    """Pure rule: whether a run with this formatting can't be seen. `color` and
    `background` are hex RGB ("FFFFFF") or None when unset; a run with no
    colour set is the document's default (black), and no background means the
    white page."""
    if vanished:
        return True
    if size_pt is not None and size_pt < _MIN_VISIBLE_FONT_SIZE:
        return True
    return _is_near_white(color) and (background is None or _is_near_white(background))


def _is_near_white(hex_color: str | None) -> bool:
    if not hex_color or len(hex_color) != 6:
        return False
    try:
        channels = [int(hex_color[i : i + 2], 16) / 255 for i in (0, 2, 4)]
    except ValueError:
        return False
    return all(channel >= 1 - _NEAR_WHITE_TOLERANCE for channel in channels)


def _resolve(run, paragraph, read):
    """The first value `read(font)` gives, from the run's own formatting, its
    character style, then the paragraph style and the styles it is based on."""
    value = read(run.font)
    if value is not None:
        return value
    styles = []
    if run.style is not None:
        styles.append(run.style)
    style = paragraph.style
    while style is not None:
        styles.append(style)
        style = style.base_style
    for style in styles:
        value = read(style.font)
        if value is not None:
            return value
    return None


def _font_color(font) -> str | None:
    color = font.color
    if color is None or color.type is None:
        return None
    if color.theme_color is not None and color.theme_color.name in _WHITE_THEME_COLORS:
        return "FFFFFF"
    return str(color.rgb) if color.rgb is not None else None


def _shading_fill(properties) -> str | None:
    """The fill colour of a w:shd under a run, paragraph or cell's properties,
    or None when there is none, it is "auto", or the properties are absent."""
    if properties is None:
        return None
    from docx.oxml.ns import qn

    shading = properties.find(qn("w:shd"))
    if shading is None:
        return None
    fill = shading.get(qn("w:fill"))
    return None if not fill or fill.lower() == "auto" else fill.upper()


# Highlight colours (w:highlight) as hex, for the white-text check. Only the
# light ones matter for "is the background white"; any other highlight is a
# coloured background, given as black so white text on it counts as visible.
_LIGHT_HIGHLIGHTS = {"white": "FFFFFF"}


def _highlight(properties) -> str | None:
    if properties is None:
        return None
    from docx.oxml.ns import qn

    highlight = properties.find(qn("w:highlight"))
    if highlight is None:
        return None
    value = (highlight.get(qn("w:val")) or "").lower()
    if not value or value == "none":
        return None
    return _LIGHT_HIGHLIGHTS.get(value, "000000")


def _lines_to_text(lines: Iterable[str]) -> str:
    """Joins lines with newlines, dropping blank runs so section headings the
    sectioner looks for ("Experience" alone on a line) stay on their own
    line while empty spacer paragraphs don't pad the word count."""
    kept = [line.strip() for line in lines]
    return "\n".join(line for line in kept if line)
